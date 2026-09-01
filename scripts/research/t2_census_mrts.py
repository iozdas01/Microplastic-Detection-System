"""Tier 2 — Census MRTS: monthly furniture retail sales (NAICS 442 / 4421).

The MRTS API endpoint is behind a Census key; the published workbook is not, and it
carries the same numbers back to 1992. One sheet per year, a NOT ADJUSTED block and an
ADJUSTED block, months across the columns.

Answers: is furniture demand currently rising or falling, and by how much.
"""
from __future__ import annotations

import re
import sys

import openpyxl

sys.path.insert(0, str(__import__("pathlib").Path(__file__).parent))
from common import RAW, guard, log, record, save_csv  # noqa: E402

URL = "https://www.census.gov/retail/mrts/www/mrtssales92-present.xlsx"
SOURCE = "census-mrts"
WANTED = {"442": "Furniture and home furnishings stores",
          "4421": "Furniture stores",
          "442,443": "Furniture, home furn, electronics and appliance stores"}
MONTHS = {m: i + 1 for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])}


def download() -> "pathlib.Path":  # noqa: F821
    import requests
    d = RAW / SOURCE
    d.mkdir(parents=True, exist_ok=True)
    p = d / "mrtssales92-present.xlsx"
    r = requests.get(URL, timeout=120,
                     headers={"User-Agent": "furniture-market-research/1.0"})
    r.raise_for_status()
    p.write_bytes(r.content)
    return p


def parse_month(label: str) -> tuple[int, int] | None:
    m = re.match(r"\s*([A-Z][a-z]{2})\.?\s+(\d{4})", str(label or ""))
    if not m or m.group(1) not in MONTHS:
        return None
    return int(m.group(2)), MONTHS[m.group(1)]


def to_num(v):
    """MRTS uses (S) suppressed, (NA) not available; keep those out of the series."""
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if not s or s.startswith("("):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def main() -> None:
    with guard(SOURCE, "monthly-retail-442", URL):
        path = download()
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for sheet in wb.sheetnames:
            if not re.fullmatch(r"\d{4}", sheet):
                continue
            rows = list(wb[sheet].iter_rows(max_col=20, values_only=True))
            # Column -> (year, month), read off the one header row that carries dates.
            colmap: dict[int, tuple[int, int]] = {}
            for r in rows[:8]:
                for ci, cell in enumerate(r):
                    ym = parse_month(cell)
                    if ym:
                        colmap[ci] = ym
                if colmap:
                    break
            adjusted = False
            for r in rows:
                label = str(r[1]).strip() if r[1] else ""
                if label.upper().startswith("NOT ADJUSTED"):
                    adjusted = False
                elif label.upper().startswith("ADJUSTED"):
                    adjusted = True
                code = str(r[0]).strip() if r[0] else ""
                if code not in WANTED:
                    continue
                for ci, (yr, mo) in colmap.items():
                    val = to_num(r[ci]) if ci < len(r) else None
                    if val is None:
                        continue
                    out.append({
                        "naics": code,
                        "kind_of_business": WANTED[code],
                        "year": yr,
                        "month": mo,
                        "period": f"{yr}-{mo:02d}",
                        "sales_musd": val,
                        "seasonally_adjusted": adjusted,
                    })
        out.sort(key=lambda r: (r["naics"], r["seasonally_adjusted"], r["period"]))
        save_csv("mrts_furniture_monthly", out)
        record(SOURCE, "monthly-retail-442", URL, len(out))
        log(f"  MRTS: {len(out)} observations, "
            f"{min(r['period'] for r in out)} → {max(r['period'] for r in out)}")


if __name__ == "__main__":
    main()
